from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart # 차트 사용을 위한 import
# 구글에서 openpyxl 검색으로 자세한 공부가능

wb = load_workbook("sample.xlsx")
ws = wb.active

# B2:C11 까지의 데이터를 차트로 생성
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3) # (ws 액티브 시트, 시작row, 끝row, 시작col, 끝col)
# bar_chart = BarChart() # 차트 종류 설정 (Bar, Line, Pie...)
# bar_chart.add_data(bar_value) # 차트 데이터 추가

# ws.add_chart(bar_chart, "E1") # 차트 넣을 위치 정의
#                               # 번호, 영어, 수학, , 차트


# B1:C11 까지의 데이터 (index 부분까지 포함)
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # 계열 > 영어, 수학()
line_chart.title = "성적표"
line_chart.style = 10 # 미리 정의된 스타일을 적용, 사용자가 개별 지정 가능
line_chart.y_axis.title = "점수" # Y축의 제목
line_chart.x_axis.title = "번호" # X축의 제목
ws.add_chart(line_chart,"E1")

wb.save("sample_chart.xlsx") 